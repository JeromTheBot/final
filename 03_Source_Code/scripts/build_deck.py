"""Build 08_Presentation.pptx (10-minute panel presentation) from live project results.

    python -m scripts.build_deck
"""
from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

SRC = Path(__file__).resolve().parent.parent
SUB = SRC.parent
ART, RES = SRC / "data" / "artifacts", SRC / "evaluation" / "results"

NAVY, TEAL, AMBER = RGBColor(0x0B, 0x1F, 0x3A), RGBColor(0x0E, 0x9A, 0xA7), RGBColor(0xF2, 0xA5, 0x41)
INK, MUTED, CARD, WHITE = RGBColor(0x1B, 0x26, 0x35), RGBColor(0x5B, 0x6B, 0x7F), RGBColor(0xEE, 0xF3, 0xF8), RGBColor(255, 255, 255)
ICE, RED = RGBColor(0xCF, 0xE8, 0xEC), RGBColor(0xC0, 0x39, 0x2B)
HEAD, BODY = "Cambria", "Calibri"


def load(p: Path, default=None):
    return json.loads(p.read_text()) if p.exists() else default


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.blank = self.prs.slide_layouts[6]
        self.n = 0

    # ---- primitives -------------------------------------------------------------------------------
    def slide(self, dark=False, notes=""):
        s = self.prs.slides.add_slide(self.blank)
        self.n += 1
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = NAVY if dark else WHITE
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        if self.n > 1:
            self.text(s, 12.2, 7.0, 0.8, 0.3, str(self.n), 10, MUTED if not dark else ICE, align=PP_ALIGN.RIGHT)
        return s

    def text(self, s, x, y, w, h, txt, size=16, color=INK, bold=False, font=BODY, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, italic=False, name=None):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            tb.name = name
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.04)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        tf.vertical_anchor = anchor
        lines = txt if isinstance(txt, list) else [txt]
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            runs = ln if isinstance(ln, list) else [(ln, {})]
            for rt, opt in runs:
                r = p.add_run()
                r.text = rt
                f = r.font
                f.size, f.name = Pt(opt.get("size", size)), opt.get("font", font)
                f.bold, f.italic = opt.get("bold", bold), opt.get("italic", italic)
                f.color.rgb = opt.get("color", color)
            p.space_after = Pt(opt.get("after", 4) if runs else 4)
        return tb

    def bullets(self, s, x, y, w, h, items, size=15, color=INK):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            head, rest = (it if isinstance(it, tuple) else ("", it))
            r0 = p.add_run()
            r0.text = "▸  "
            r0.font.size, r0.font.color.rgb, r0.font.name = Pt(size), TEAL, BODY
            if head:
                r = p.add_run()
                r.text = head + " "
                r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), True, color, BODY
            r = p.add_run()
            r.text = rest
            r.font.size, r.font.color.rgb, r.font.name = Pt(size), color, BODY
            p.space_after = Pt(9)
        return tb

    def box(self, s, x, y, w, h, fill=CARD, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
        sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        if line is None:
            sh.line.fill.background()
        else:
            sh.line.color.rgb = line
            sh.line.width = Pt(1.25)
        sh.shadow.inherit = False
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = radius
        return sh

    def circle(self, s, x, y, d, label, fill=TEAL, size=16, color=WHITE):
        c = self.box(s, x, y, d, d, fill, shape=MSO_SHAPE.OVAL)
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        r = p.add_run()
        r.text = label
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), True, color, BODY
        return c

    def title(self, s, t, sub=None, dark=False):
        self.text(s, 0.6, 0.38, 12.1, 0.8, t, 34, WHITE if dark else NAVY, bold=True, font=HEAD, name="Title")
        if sub:
            self.text(s, 0.62, 1.12, 12.1, 0.45, sub, 16, ICE if dark else MUTED, italic=True)

    def stat(self, s, x, y, w, value, label, vcolor=TEAL, bg=CARD, lcolor=INK):
        self.box(s, x, y, w, 1.75, bg)
        self.text(s, x + 0.15, y + 0.15, w - 0.3, 0.9, value, 40, vcolor, bold=True, font=HEAD, align=PP_ALIGN.CENTER)
        self.text(s, x + 0.15, y + 1.05, w - 0.3, 0.6, label, 13, lcolor, align=PP_ALIGN.CENTER)

    def bar_chart(self, s, x, y, w, h, cats, series: dict, title, colors=(TEAL, AMBER, NAVY), pct=False, maxv=None):
        cd = CategoryChartData()
        cd.categories = cats
        for k, v in series.items():
            cd.add_series(k, v)
        gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        ch = gf.chart
        ch.has_title = True
        ch.chart_title.text_frame.text = title
        tp = ch.chart_title.text_frame.paragraphs[0].runs[0].font
        tp.size, tp.bold, tp.color.rgb, tp.name = Pt(13), True, INK, BODY
        ch.has_legend = len(series) > 1
        if ch.has_legend:
            ch.legend.include_in_layout = False
            ch.legend.font.size = Pt(11)
            ch.legend.font.color.rgb = MUTED
        va, ca = ch.value_axis, ch.category_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = RGBColor(0xDD, 0xE3, 0xEA)
        va.tick_labels.font.size = Pt(10)
        va.tick_labels.font.color.rgb = MUTED
        va.format.line.fill.background()
        if maxv:
            va.maximum_scale = maxv
        va.minimum_scale = 0
        ca.tick_labels.font.size = Pt(11)
        ca.tick_labels.font.color.rgb = INK
        ca.format.line.color.rgb = RGBColor(0xC8, 0xD0, 0xDA)
        ca.reverse_order = True
        for i, ser in enumerate(ch.plots[0].series):
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = colors[i % len(colors)]
        pl = ch.plots[0]
        pl.gap_width = 60
        pl.has_data_labels = True
        dl = pl.data_labels
        dl.font.size, dl.font.color.rgb = Pt(10), INK
        dl.number_format = '0%' if pct else '0.00'
        dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        return ch


def main():
    met = load(ART / "model_metrics.json", {})
    rec = load(RES / "recommendation_eval.json", {}).get("summary", {})
    ret = load(RES / "retrieval_eval.json", {})
    meta = load(ART / "index_meta.json", {})
    n_cases = meta.get("n_cases", 100000)
    d = Deck()

    # 1 ---------------------------------------------------------------- title
    s = d.slide(dark=True, notes="Introduce yourself and the project: Project 4 of the PU Bootcamp GenAI capstone. "
                                 "One sentence: we turn 100K historical 5G network records into an assistant that "
                                 "finds similar incidents, explains them and recommends validated fixes.")
    d.circle(s, 0.75, 1.0, 1.1, "📡", fill=TEAL, size=30)
    d.text(s, 0.75, 2.4, 11.5, 1.8, "AI-Powered Telecom Network Operations Assistant", 46, WHITE, bold=True, font=HEAD)
    d.text(s, 0.78, 4.25, 12.0, 0.6, "Hybrid RAG · multi-agent troubleshooting · fault-risk prediction · LLM-as-judge",
           20, ICE)
    d.text(s, 0.78, 5.7, 11, 0.9, [[("Jerom Santhiyagu", {"bold": True, "color": WHITE, "size": 18})],
                                   [("PU Bootcamp GenAI Capstone — Project 4 · Prodapt", {"color": ICE, "size": 14})]], 16)

    # 2 ---------------------------------------------------------------- problem
    s = d.slide(notes="Engineers in the NOC deal with recurring congestion, latency, call drops and faults. "
                      "Searching history manually is slow. These three questions come straight from the brief "
                      "and are what the system is built to answer.")
    d.title(s, "NOC engineers lose time searching history", "Thousands of records, recurring problems, no fast way to reuse past fixes")
    d.bullets(s, 0.6, 1.9, 5.3, 4.8, [
        ("Recurring issues:", "congestion, poor signal, high latency, packet loss, call drops, handover failures, faults"),
        ("Rich history:", "KPIs, alarms, config changes, users, security events, weather"),
        ("But:", "finding similar incidents, causes and corrective actions by hand is slow and inconsistent"),
        ("Goal:", "ask in plain English → similar cases, analysis, risk score and recommended actions")], 16)
    qs = ["A cell is experiencing high latency, packet loss and reduced throughput. What similar conditions occurred before?",
          "Multiple cells show poor service quality and increasing alarm counts. What factors could contribute?",
          "A region has increasing fault occurrence and declining handover success. What corrective actions?"]
    for i, q in enumerate(qs):
        y = 1.9 + i * 1.6
        d.box(s, 6.4, y, 6.3, 1.35, CARD)
        d.circle(s, 6.6, y + 0.35, 0.65, "?", fill=AMBER, size=22, color=NAVY)
        d.text(s, 7.45, y + 0.12, 5.1, 1.15, f"“{q}”", 15, INK, italic=True, anchor=MSO_ANCHOR.MIDDLE)

    # 3 ---------------------------------------------------------------- at a glance
    s = d.slide(notes="The solution at a glance. Every number here is produced by the code in the repo.")
    d.title(s, "The solution at a glance", "One microservice, one question, one grounded and validated answer")
    stats = [(f"{n_cases // 1000}K", "incidents indexed\nFAISS + BM25"), ("7", "agents with a\njudge loop"),
             ("13", "REST endpoints\nFastAPI + Swagger"),
             (f"{rec.get('category_correct', 1):.0%}", "category accuracy\ngolden queries"),
             (f"{rec.get('latency_ms', 650) / 1000:.1f}s", "mean analysis time\noffline mode")]
    for i, (v, l) in enumerate(stats):
        d.stat(s, 0.6 + i * 2.48, 1.95, 2.28, v, l)
    d.text(s, 0.6, 4.1, 12.1, 0.4, "What an engineer gets back", 20, NAVY, bold=True, font=HEAD)
    outs = [("🔎", "Similar cases", "hybrid search + filters"), ("🏷️", "Priority P1–P4", "category + confidence"),
            ("📈", "Forecasts", "fault risk + KPIs"), ("🔬", "Root cause", "impact, SLA, MTTR"),
            ("🛠️", "Action plan", "cited, approval-flagged"), ("⚖️", "Judge verdict", "5 quality criteria")]
    for i, (ic, h, t) in enumerate(outs):
        x = 0.6 + i * 2.07
        d.circle(s, x + 0.62, 4.7, 0.7, ic, size=20)
        d.text(s, x, 5.5, 1.95, 0.4, h, 15, INK, bold=True, align=PP_ALIGN.CENTER)
        d.text(s, x, 5.9, 1.95, 0.5, t, 12, MUTED, align=PP_ALIGN.CENTER)

    # 4 ---------------------------------------------------------------- architecture
    s = d.slide(notes="Walk the diagram left to right. Top: the offline pipeline that turns the CSV into indexes and "
                      "models. Middle: the online path through FastAPI, guardrails and the seven agents, with the judge "
                      "able to send the plan back once. Bottom: LLM, feedback, monitoring and evaluation.")
    d.title(s, "Architecture: build offline, answer online")
    img = SUB / "01_Architecture.jpeg"
    if img.exists():
        h = 5.9
        w = h * 2420 / 1452
        s.shapes.add_picture(str(img), Inches((13.333 - w) / 2), Inches(1.25), Inches(w), Inches(h))

    # 5 ---------------------------------------------------------------- data
    s = d.slide(notes="The dataset is numbers only, so the key design step is turning KPIs into the language "
                      "engineers use. Thresholds are percentiles of the fleet, so they adapt to the data. Mention the "
                      "two honest findings: the fault rate is near-deterministic and Issue Reported is noise.")
    d.title(s, "From 100K KPI rows to searchable incidents", "5G NetOps dataset (Hugging Face fenar/5g_netops) · 27 fields · one record = one chunk")
    steps = [("1", "Ingest & clean", "download, dedupe, impute, clip bounds, fix encoding"),
             ("2", "21 KPI conditions", "degraded ≥ fleet P75, severe ≥ P90 (data-driven)"),
             ("3", "Taxonomy & risk", "8 incident classes · risk 0–100 · priority P1–P4"),
             ("4", "Narrative + index", "summary → MiniLM/FAISS · keywords → BM25")]
    for i, (n, h, t) in enumerate(steps):
        x = 0.6 + i * 3.1
        d.box(s, x, 1.95, 2.85, 2.1, CARD)
        d.circle(s, x + 0.2, 2.15, 0.6, n, size=18)
        d.text(s, x + 0.95, 2.22, 1.85, 0.5, h, 16, NAVY, bold=True)
        d.text(s, x + 0.2, 2.95, 2.5, 1.0, t, 13, INK)
    d.box(s, 0.6, 4.4, 12.1, 2.3, RGBColor(0xFD, 0xF3, 0xE4))
    d.text(s, 0.85, 4.55, 11.6, 0.4, "Data findings that shaped the design", 17, NAVY, bold=True, font=HEAD)
    fr = met.get("fault_risk", {})
    d.bullets(s, 0.85, 5.05, 11.6, 1.6, [
        ("Fault rate is near-deterministic", f"(R² {fr.get('r2', 1.0):.3f}; critical alarms r = 0.93) → added an early-warning "
                                             f"model on leading indicators (R² {fr.get('early_warning_r2', 0.64):.2f})"),
        ("'Issue Reported' is label noise", "(identical KPI means for yes/no) → used only as a filter"),
        ("No free text", "→ narratives generated from conditions; numbers kept out of BM25 to avoid vocabulary noise")], 14)

    # 6 ---------------------------------------------------------------- retrieval
    st = ret.get("strategies", {})
    s = d.slide(notes="Hybrid retrieval: BM25 on a keyword field plus dense MiniLM vectors in FAISS, fused with "
                      "Reciprocal Rank Fusion and re-ranked by symptom overlap and feedback. Metadata filters are "
                      "applied before search, so top-k is always inside the filter. The chart compares strategies on "
                      "paraphrased queries.")
    d.title(s, "Hybrid + symptom re-rank wins retrieval", "Fusion alone does not beat BM25; re-ranking the fused pool on symptoms does")
    d.bullets(s, 0.6, 1.9, 5.4, 4.9, [
        ("Dense:", "all-MiniLM-L6-v2 → FAISS exact cosine"),
        ("Sparse:", "vectorised BM25 (uni+bi-grams) on location, season, weather, conditions, category"),
        ("Fusion:", "Reciprocal Rank Fusion, semantic weight α = 0.6 (tunable per request)"),
        ("Re-rank:", "+ symptom overlap with query, + feedback boost. This is where the gain comes from"),
        ("Filters:", "state, city, season, weather, category, priority, risk, required conditions; auto-detected from text")], 15)
    if st:
        names = {"keyword_bm25": "BM25 only", "semantic_dense": "Semantic only", "hybrid_rrf": "Hybrid (RRF)",
                 "hybrid_rrf+symptom_rerank": "Hybrid + re-rank"}
        cats = [names[k] for k in names if k in st]
        d.bar_chart(s, 6.3, 1.8, 6.5, 5.0, cats, {"nDCG@10": [st[k]["ndcg@10"] for k in names if k in st],
                                                    "P@5": [st[k]["p@5"] for k in names if k in st]},
                    "Retrieval quality, 100K index, 40 paraphrased queries", maxv=1.0)

    # 7 ---------------------------------------------------------------- agents
    s = d.slide(notes="Seven agents share one context object. Each leaves a trace entry. The judge scores the plan on "
                      "five criteria; if it says revise, its issues go back to the recommendation agent once. Every "
                      "agent has a deterministic fallback, so the system works without an LLM.")
    d.title(s, "Seven agents collaborate on every analysis", "Supervisor + shared blackboard · full trace · deterministic fallback for every LLM step")
    ag = [("1", "Query understanding", "symptoms, severity, location"), ("2", "Retrieval", "hybrid RAG + filters"),
          ("3", "Classification", "rules + k-NN vote (+LLM)"), ("4", "Risk prediction", "case-based KPI profile → ML"),
          ("5", "Root cause & impact", "deviations, drivers, users, SLA"), ("6", "Recommendation", "SOP RAG, cited plan"),
          ("7", "LLM-as-judge", "5 criteria → pass / revise")]
    for i, (n, h, t) in enumerate(ag):
        x = 0.6 + i * 1.77
        d.circle(s, x + 0.45, 2.1, 0.8, n, fill=AMBER if n == "7" else TEAL, size=22, color=NAVY if n == "7" else WHITE)
        d.text(s, x, 3.05, 1.7, 0.75, h, 13, NAVY, bold=True, align=PP_ALIGN.CENTER)
        d.text(s, x, 3.75, 1.7, 0.8, t, 12, MUTED, align=PP_ALIGN.CENTER)
        if i < 6:
            arr = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 1.4), Inches(2.4), Inches(0.28), Inches(0.22))
            arr.fill.solid()
            arr.fill.fore_color.rgb = RGBColor(0xB5, 0xC2, 0xD0)
            arr.line.fill.background()
    d.box(s, 0.6, 4.85, 12.1, 1.85, CARD)
    d.text(s, 0.85, 4.98, 11.6, 0.4, "Collaboration loop", 17, NAVY, bold=True, font=HEAD)
    d.bullets(s, 0.85, 5.45, 11.6, 1.2, [
        ("Judge → Recommendation:", "if groundedness ≤ 3 or any score ≤ 2, the issues are sent back and the plan is rebuilt "
                                    "(e.g. adds SOPs for uncovered symptoms); at most one revision bounds cost"),
        ("Why plain Python, not LangGraph:", "a fixed DAG plus one loop gives deterministic, testable, fast control flow")], 14)

    # 8 ---------------------------------------------------------------- prediction & RCA
    s = d.slide(notes="Fault-risk models: the full model is used for what-if analysis, and the early-warning model is "
                      "the honest forecast. The performance models beat the baseline but are modest, and we show R² "
                      "in the UI. Root cause triangulates statistics, model drivers and history.")
    d.title(s, "Predicting risk, explaining the root cause", "HistGradientBoosting models with local driver explanations")
    pf = met.get("performance", {})
    cats = ["Fault risk (full)", "Fault early-warning"] + [f"Perf: {k.split(' (')[0]}" for k in pf]
    vals = [fr.get("r2", 0), fr.get("early_warning_r2", 0)] + [v["r2"] for v in pf.values()]
    d.bar_chart(s, 0.5, 1.75, 6.3, 4.9, cats, {"R² (20K test rows)": vals}, "Model accuracy (R², held-out 20%)",
                colors=(TEAL,), maxv=1.0)
    d.text(s, 7.2, 1.85, 5.5, 0.4, "Root cause = three independent signals", 17, NAVY, bold=True, font=HEAD)
    sig = [("σ", "KPI deviation", "direction-aware z-scores vs fleet baseline"),
           ("ƒ", "Model drivers", "median-substitution ablation on the fault model"),
           ("≈", "Historical evidence", "condition prevalence in the top-k similar cases")]
    for i, (ic, h, t) in enumerate(sig):
        y = 2.45 + i * 1.05
        d.circle(s, 7.2, y, 0.7, ic, size=20)
        d.text(s, 8.05, y - 0.02, 4.7, 0.4, h, 15, INK, bold=True)
        d.text(s, 8.05, y + 0.35, 4.7, 0.5, t, 13, MUTED)
    d.box(s, 7.2, 5.65, 5.5, 1.05, CARD)
    d.text(s, 7.4, 5.75, 5.1, 0.9, [[("Impact: ", {"bold": True}), ("users affected, data at risk, availability gap, "
                                                                    "expected MTTR, SLA breach risk, services affected, "
                                                                    "pattern prevalence", {})]], 13)

    # 9 ---------------------------------------------------------------- guardrails & feedback
    s = d.slide(notes="Security and compliance: validate everything at the edge, redact PII before any LLM call, and "
                      "sanitise outputs. The feedback loop makes the system better with use, and a test proves the "
                      "ranking changes.")
    d.title(s, "Safe by design, better with every rating", "Guardrails on input and output · feedback flows back into ranking and actions")
    cols = [("🛡️", "Input guardrails", ["Pydantic ranges & enums on every field", "13 prompt-injection / SQL / script patterns → 422",
                                       "PII redaction: email, phone, IMSI, card", "domain-relevance check (off-topic → 422)"]),
            ("✅", "Output guardrails", ["citations must be retrieved case IDs or SOPs", "unsafe steps removed (e.g. disable alarms)",
                                        "service-affecting steps need approval", "optional X-API-Key, request-ID audit trail"]),
            ("🔁", "Feedback loop", ["1–5 rating, helpful cases, effective actions", "case boost/penalty in re-ranking (±0.15)",
                                    "actions re-ordered per category", "category corrections → new labelled examples"])]
    for i, (ic, h, items) in enumerate(cols):
        x = 0.6 + i * 4.1
        d.box(s, x, 1.9, 3.85, 4.8, CARD)
        d.circle(s, x + 0.25, 2.1, 0.75, ic, size=22)
        d.text(s, x + 1.15, 2.25, 2.6, 0.5, h, 18, NAVY, bold=True, font=HEAD)
        d.bullets(s, x + 0.25, 3.1, 3.4, 3.5, items, 14)

    # 10 --------------------------------------------------------------- evaluation
    s = d.slide(notes="Evaluation at three levels: unit, API and agent tests; retrieval metrics; and recommendation "
                      "quality on golden queries with the heuristic judge and, when a key is set, DeepEval metrics. "
                      "Note the embedding comparison that justified MiniLM.")
    d.title(s, "Measured, not assumed", "pytest suite · retrieval benchmark · golden-set evaluation · DeepEval when an LLM is configured")
    ev = [(f"{rec.get('condition_f1', 0):.2f}", "symptom-extraction F1"), (f"{rec.get('sop_coverage', 0):.0%}", "SOP coverage"),
          (f"{rec.get('citation_validity', 0):.0%}", "valid citations"), (f"{rec.get('judge_overall', 0):.1f}/5", "judge score")]
    for i, (v, l) in enumerate(ev):
        d.stat(s, 0.6 + (i % 2) * 2.75, 1.9 + (i // 2) * 2.0, 2.55, v, l)
    em = ret.get("embedding_models", {})
    if em:
        cats = [m.split("/")[-1] for m in em]
        d.bar_chart(s, 6.3, 1.75, 6.5, 4.4, cats, {"dense nDCG@10": [v["dense_only"]["ndcg@10"] for v in em.values()],
                                                    "hybrid nDCG@10": [v["hybrid+rerank"]["ndcg@10"] for v in em.values()]},
                    "Embedding models (4K-case subset)", maxv=1.0)
    d.text(s, 0.6, 6.05, 12.1, 0.7, [[("DeepEval metrics: ", {"bold": True, "color": NAVY}),
                                      ("AnswerRelevancy · Faithfulness · ContextualRelevancy · GEval Recommendation "
                                       "Quality · GEval Diagnosis Correctness (judge = configured LLM)", {"color": INK})]], 14)

    # 11 --------------------------------------------------------------- demo
    s = d.slide(notes="Switch to the live demo now and follow 07_Demo_Script.md.")
    d.title(s, "Live demo", "Streamlit UI on top of the FastAPI microservice")
    flow = [("1", "Ask", "brief question 1 → full multi-agent analysis"), ("2", "Filter", "hybrid vs keyword vs semantic; Winter, P1/P2, TX"),
            ("3", "Guardrails", "prompt injection blocked, PII redacted"), ("4", "Prioritise", "batch of incidents ranked P1 → P4"),
            ("5", "What-if", "risk and performance predictor with drivers"), ("6", "Feedback", "rate → monitoring and feedback dashboard")]
    for i, (n, h, t) in enumerate(flow):
        x, y = 0.6 + (i % 3) * 4.1, 1.95 + (i // 3) * 2.45
        d.box(s, x, y, 3.85, 2.15, CARD)
        d.circle(s, x + 0.25, y + 0.25, 0.7, n, size=20)
        d.text(s, x + 1.1, y + 0.33, 2.6, 0.5, h, 19, NAVY, bold=True, font=HEAD)
        d.text(s, x + 0.25, y + 1.15, 3.4, 0.9, t, 14, INK)

    # 12 --------------------------------------------------------------- trade-offs & next steps
    s = d.slide(dark=True, notes="Close with the trade-offs we chose deliberately, and what we would do next with "
                                 "real network data. Then open for questions.")
    d.title(s, "Trade-offs we chose, and what comes next", dark=True)
    d.text(s, 0.6, 1.5, 6, 0.45, "Deliberate trade-offs", 20, AMBER, bold=True, font=HEAD)
    d.bullets(s, 0.6, 2.05, 6.0, 4.6, [
        ("Deterministic core, LLM for language:", "always available and cheap; less fluent offline"),
        ("Exact FAISS at 100K:", "no recall loss; move to ANN / managed vector DB beyond about 5M vectors"),
        ("Small LLMs (4o-mini / Flash / Haiku):", "evidence-bound prompts do not need large models"),
        ("Percentile thresholds:", "adapt to the fleet, but are relative rather than 3GPP targets")], 15, WHITE)
    d.text(s, 7.0, 1.5, 5.8, 0.45, "Next steps", 20, AMBER, bold=True, font=HEAD)
    d.bullets(s, 7.0, 2.05, 5.7, 4.6, [
        ("Real data:", "retrain on live KPIs with timestamps and cell topology"),
        ("Retrieval:", "cross-encoder re-ranker, LLM query rewriting"),
        ("Ops:", "Azure Container Apps, Azure AI Search, Prometheus / Azure Monitor"),
        ("Learning:", "A/B tests on feedback ranking, retrain the classifier on corrections")], 15, WHITE)
    d.text(s, 0.6, 6.55, 12.1, 0.5, "Thank you — questions?", 22, WHITE, bold=True, font=HEAD, align=PP_ALIGN.CENTER)

    out = SUB / "08_Presentation.pptx"
    d.prs.save(out)
    print("written", out, d.n, "slides")


if __name__ == "__main__":
    main()
